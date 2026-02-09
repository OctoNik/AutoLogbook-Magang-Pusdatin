
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

def create_template():
    doc = Document()

    # --- 1. SETUP FONT DEFAULT (Times New Roman) ---
    style = doc.styles['Normal']
    style.font.name = 'Times New Roman'
    style.font.size = Pt(12)
    
    # --- 2. HEADER/JUDUL ---
    # Logbook Dynamic Title
    p_title = doc.add_paragraph('Logbook {{ judul_logbook }}')
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.style.font.bold = True
    
    # Richie Lagito
    p_name = doc.add_paragraph('Richie Lagito')
    p_name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_name.style.font.bold = True
    
    # Spasi
    doc.add_paragraph('')

    # --- 3. TABEL (HANYA HEADER) ---
    # Membuat tabel 1 baris (Header saja). Data akan diisi via Python-docx manual.
    table = doc.add_table(rows=1, cols=4)
    table.style = 'Table Grid'
    table.autofit = False 
    
    # Atur Lebar Kolom (Estimasi agar rapi)
    widths = [Cm(1), Cm(4), Cm(6), Cm(5)]
    for row in table.rows:
        for idx, width in enumerate(widths):
            row.cells[idx].width = width

    # -- Isi Header --
    hdr_cells = table.rows[0].cells
    headers = ['No', 'Hari dan Tanggal', 'Kegiatan', 'Foto']
    for idx, text in enumerate(headers):
        hdr_cells[idx].text = text
        hdr_cells[idx].paragraphs[0].runs[0].font.bold = True
        hdr_cells[idx].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    # --- 4. FOOTER / TANDA TANGAN ---
    doc.add_paragraph('\n') # Spasi kosong
    
    # Tempat & Tanggal (Dibuat variabel agar bisa ganti bulan otomatis)
    p_date = doc.add_paragraph('Jakarta, {{ tanggal_ttd }}')
    p_date.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    
    # Mengetahui
    p_sign = doc.add_paragraph('Mengetahui,')
    p_sign.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    
    # Spasi untuk tanda tangan
    doc.add_paragraph('\n\n\n')
    
    # Nama Dosen Pembimbing
    p_dosen = doc.add_paragraph('Ahmad Hulaliyah S.Kom., M.Kom.')
    p_dosen.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    # Garis bawah nama (opsional, sesuaikan kebiasaan)
    p_dosen.runs[0].font.underline = True

    # --- SAVE ---
    output_filename = 'template_logbook.docx'
    doc.save(output_filename)
    print(f"Berhasil! File '{output_filename}' sudah dibuat.")
    print("Silahkan buka file tersebut untuk mengecek layoutnya.")

if __name__ == "__main__":
    create_template()